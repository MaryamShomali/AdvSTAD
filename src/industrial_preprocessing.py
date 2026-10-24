"""Preprocess the official SWaT physical and revised WADI releases."""

import csv
import json
from pathlib import Path

import numpy as np
import pandas as pd


WADI_LABEL = "Attack LABLE (1:No Attack, -1:Attack)"
SWAT_LABEL = "Normal/Attack"


def _clean_columns(columns, source):
    names = [str(name).strip() for name in columns]
    if len(names) != len(set(names)):
        raise ValueError(f"Duplicate column names in {source}")
    return names


def _read_wadi(path, attack=False):
    # The attack release starts with a numeric index line before its header.
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for header, row in enumerate(csv.reader(handle)):
            if {"Row", "Date", "Time"}.issubset(name.strip() for name in row):
                break
            if header >= 4:
                raise ValueError(f"WADI header not found in {path}")
        else:
            raise ValueError(f"WADI header not found in {path}")
    frame = pd.read_csv(path, header=header, encoding="utf-8-sig", low_memory=False)
    frame.columns = _clean_columns(frame.columns, path)
    if attack and WADI_LABEL not in frame:
        raise ValueError(f"Missing {WADI_LABEL!r} in {path}")
    # Two footer rows have a label but no observation; drop them before decoding.
    observation_columns = [name for name in frame if name != WADI_LABEL]
    frame = frame.loc[~frame[observation_columns].isna().all(axis=1)]
    if frame[["Row", "Date", "Time"]].isna().any().any():
        raise ValueError(f"Missing WADI observation metadata in {path}")
    labels = None
    if attack:
        raw_labels = pd.to_numeric(frame[WADI_LABEL], errors="raise")
        if not raw_labels.isin([1, -1]).all():
            raise ValueError(f"Unknown WADI attack label in {path}; expected 1 or -1")
        labels = (raw_labels.to_numpy() == -1).astype(np.uint8)
    features = frame.drop(columns=["Row", "Date", "Time", WADI_LABEL], errors="ignore")
    return features, labels


def _read_swat(path, sheet_name, normal=False):
    # Stream cells instead of retaining half a million rows of Python objects.
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        for header_row, row in enumerate(rows, start=1):
            columns = [str(name).strip() for name in row]
            if "Timestamp" in columns and SWAT_LABEL in columns:
                columns = _clean_columns(row, path)
                break
            if header_row >= 5:
                raise ValueError(f"SWaT header not found in {path}")
        else:
            raise ValueError(f"SWaT header not found in {path}")
        feature_indices = [i for i, name in enumerate(columns)
                           if name not in ("Timestamp", SWAT_LABEL)]
        names = [columns[i] for i in feature_indices]
        timestamp_index = columns.index("Timestamp")
        label_index = columns.index(SWAT_LABEL)
        values = np.empty((sheet.max_row - header_row, len(names)), dtype=np.float64)
        labels = np.empty(values.shape[0], dtype=np.uint8)
        count = 0
        for source_row, row in enumerate(rows, start=header_row + 1):
            if row[timestamp_index] is None and all(row[i] is None for i in feature_indices):
                continue
            if row[timestamp_index] is None:
                raise ValueError(f"Missing SWaT timestamp in {path}, row {source_row}")
            # The attack workbook also contains the spelling 'A ttack'.
            label = "".join(str(row[label_index]).split()).lower()
            if label not in ("normal", "attack"):
                raise ValueError(f"Unknown SWaT label {row[label_index]!r} in {path}, row {source_row}")
            if normal and label != "normal":
                raise ValueError(f"Attack observation in SWaT normal data: {path}, row {source_row}")
            try:
                values[count] = [np.nan if row[i] is None else float(row[i]) for i in feature_indices]
            except (TypeError, ValueError) as error:
                raise ValueError(f"Non-numeric SWaT sensor in {path}, row {source_row}") from error
            labels[count] = label == "attack"
            count += 1
        return pd.DataFrame(values[:count], columns=names), labels[:count]
    finally:
        workbook.close()


def _prepare_features(train, test):
    if train.empty or test.empty:
        raise ValueError("Training and test observations must be non-empty")
    if set(train.columns) != set(test.columns):
        raise ValueError(
            "Train/test sensor columns disagree: "
            f"missing in test={sorted(set(train) - set(test))}, "
            f"extra in test={sorted(set(test) - set(train))}"
        )
    test = test.loc[:, train.columns]
    train = train.apply(pd.to_numeric, errors="raise").replace([np.inf, -np.inf], np.nan)
    test = test.apply(pd.to_numeric, errors="raise").replace([np.inf, -np.inf], np.nan)
    # Feature selection and imputation statistics use normal training data only.
    dropped = train.columns[train.isna().all()].tolist()
    train = train.drop(columns=dropped)
    test = test.drop(columns=dropped)
    if train.shape[1] == 0:
        raise ValueError("No observed training sensor columns remain")
    medians = train.median()
    train = train.ffill().fillna(medians).to_numpy(dtype=np.float64)
    test = test.ffill().fillna(medians).to_numpy(dtype=np.float64)
    minimum = train.min(axis=0)
    span = train.max(axis=0) - minimum
    scale = np.where(span == 0, 1.0, span)
    train = (train - minimum) / scale
    test = (test - minimum) / scale
    if not np.isfinite(train).all() or not np.isfinite(test).all():
        raise ValueError("Preprocessing produced non-finite sensor values")
    metadata = {
        "feature_names": medians.index.tolist(),
        "dropped_features": dropped,
        "imputation": {"method": "forward_fill_then_training_median", "medians": medians.tolist()},
        "normalization": {"method": "training_minmax", "minimum": minimum.tolist(),
                          "scale": scale.tolist(), "clip_test": False},
    }
    return train, test, metadata


def preprocess_industrial(dataset, dataset_folder, destination, downsample=1):
    """Write [timestamps, sensors] arrays and explicit timestamp-label metadata."""
    if isinstance(downsample, bool) or not isinstance(downsample, int) or downsample < 1:
        raise ValueError("downsample must be a positive integer")
    dataset_folder, destination = Path(dataset_folder), Path(destination)
    if dataset == "WADI":
        train_path = dataset_folder / "raw" / "WADI_14days_new.csv"
        test_path = dataset_folder / "raw" / "WADI_attackdataLABLE.csv"
        print(f"Reading {train_path}", flush=True)
        train_frame, _ = _read_wadi(train_path)
        print(f"Reading {test_path}", flush=True)
        test_frame, row_labels = _read_wadi(test_path, attack=True)
        release = "WADI.A2_19 Nov 2019"
    elif dataset == "SWaT":
        train_path = dataset_folder / "raw" / "Physical" / "SWaT_Dataset_Normal_v1.xlsx"
        test_path = dataset_folder / "raw" / "Physical" / "SWaT_Dataset_Attack_v0.xlsx"
        print(f"Reading {train_path}", flush=True)
        train_frame, _ = _read_swat(train_path, "Normal.csv", normal=True)
        print(f"Reading {test_path}", flush=True)
        test_frame, row_labels = _read_swat(test_path, "Combined Data")
        release = "SWaT.A1 & A2_Dec_2015 (normal v1, attack v0)"
    else:
        raise ValueError(f"Unsupported industrial dataset: {dataset}")
    source_rows = {"train": len(train_frame), "test": len(test_frame)}
    source_attacks = int(row_labels.sum())
    train, test, metadata = _prepare_features(train_frame, test_frame)
    del train_frame, test_frame
    train, test = train[::downsample], test[::downsample]
    row_labels = row_labels[::downsample]
    # Existing models/plotting expect labels.shape == test.shape. These copies
    # express a global attack at that timestamp, not attacked sensor identities.
    labels = np.broadcast_to(row_labels[:, None], test.shape).copy()
    metadata.update({
        "schema_version": 1,
        "dataset": dataset,
        "release": release,
        "source_files": {"train": str(train_path), "test": str(test_path)},
        "source_rows": source_rows,
        "source_attack_timestamps": source_attacks,
        "downsample": downsample,
        "sampling": "every_nth_observation_from_row_zero",
        "label_scope": "timestamp",
        "labels": "0=normal, 1=attack; broadcast across sensors",
        "attack_timestamps": int(row_labels.sum()),
        "shapes": {"train": list(train.shape), "test": list(test.shape), "labels": list(labels.shape)},
    })
    destination.mkdir(parents=True, exist_ok=True)
    for name, values in (("train", train), ("test", test), ("labels", labels)):
        np.save(destination / f"{name}.npy", values)
    with (destination / "metadata.json").open("w", encoding="utf-8") as handle:
        json.dump(metadata, handle, indent=2, allow_nan=False)
        handle.write("\n")
    print(f"{dataset}: train={train.shape}, test={test.shape}, labels={labels.shape}; "
          f"attack timestamps={int(row_labels.sum())}", flush=True)
    return metadata
